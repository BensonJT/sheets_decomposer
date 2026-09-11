"""openpyxl -> unified workbook dict. Handles .xlsx from Excel or a Sheets export.

Values come from the workbook's cached results (data_only=True). A file written
by a script and never opened in Excel/Sheets has NO cached values; formulas are
still fully extracted.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .model import new_cell, new_sheet, new_workbook


def fetch(path: str, source_type: str = "xlsx", source_meta: dict | None = None) -> dict:
    p = Path(path)
    wb_f = load_workbook(p, data_only=False)
    wb_v = load_workbook(p, data_only=True)
    src = {"type": source_type, "path": str(p.resolve()), "id": p.stem}
    if source_meta:
        src.update(source_meta)
    wb = new_workbook(src, title=p.stem)

    for dn_name, dn in wb_f.defined_names.items():
        for sheet, rng in dn.destinations:
            wb["named_ranges"].append({"name": dn_name, "sheet": sheet, "range": rng.replace("$", "")})

    for idx, ws in enumerate(wb_f.worksheets):
        wsv = wb_v[ws.title]
        sh = new_sheet(
            idx, idx, ws.title, hidden=(ws.sheet_state != "visible"),
            row_count=ws.max_row, col_count=ws.max_column,
            frozen_rows=(ws.freeze_panes and ws[ws.freeze_panes].row - 1) or 0,
            frozen_cols=(ws.freeze_panes and ws[ws.freeze_panes].column - 1) or 0,
            tab_color=(ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else None),
        )
        sh["hidden_rows"] = [r for r, d in ws.row_dimensions.items() if d.hidden]
        sh["hidden_cols"] = [c for c, d in ws.column_dimensions.items() if d.hidden]
        sh["merges"] = [str(m) for m in ws.merged_cells.ranges]
        if ws.protection and ws.protection.sheet:
            sh["protected_ranges"].append({"range": "(sheet)", "description": "sheet protection", "warning_only": False})
        for dv in ws.data_validations.dataValidation:
            for rng in str(dv.sqref).split():
                sh["validations"].append({"a1": rng, "type": dv.type, "values": [dv.formula1, dv.formula2], "strict": bool(dv.showErrorMessage), "show_dropdown": bool(dv.showDropDown is not True)})
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for rule in rules:
                sh["conditional_formats"].append({"ranges": [str(rng.sqref)], "type": rule.type, "values": list(rule.formula or [])})

        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.comment and not c.hyperlink:
                    continue
                v = c.value
                vv = wsv[c.coordinate].value
                if isinstance(v, str) and v.startswith("="):
                    kind, formula = "formula", v
                elif isinstance(v, bool):
                    kind, formula = "bool", None
                elif isinstance(v, (int, float)):
                    kind, formula = "number", None
                elif isinstance(v, (datetime, date)):
                    kind, formula = "date", None
                elif isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")):
                    kind, formula = "error", None
                elif v is None:
                    kind, formula = "empty", None
                else:
                    kind, formula = "string", None
                if isinstance(vv, (datetime, date)):
                    vv = vv.isoformat()
                if kind == "formula" and vv is None:
                    vv = None  # no cached value available
                sh["cells"].append(new_cell(
                    c.row, c.column, kind, formula, vv if kind == "formula" else (v.isoformat() if isinstance(v, (datetime, date)) else v),
                    None, c.number_format if c.number_format != "General" else None,
                    c.comment.text if c.comment else None, c.hyperlink.target if c.hyperlink else None,
                ))
        wb["sheets"].append(sh)
    return wb
